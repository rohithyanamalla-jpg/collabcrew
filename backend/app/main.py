import os, hashlib, secrets, hmac, base64, json, io
from datetime import datetime, timedelta, date
from typing import Optional
from fastapi import FastAPI, Depends, HTTPException, Query, Header
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_, desc
from .database import Base, engine, get_db, SessionLocal
from .models import User, CreatorProfile, PortfolioProject, CreativeBrief, Application, Collaboration, VerificationEvidence
from .schemas import RegisterIn, LoginIn, ProfileIn, PortfolioIn, BriefIn, ApplicationIn, StatusIn, EvidenceIn
import requests
from dotenv import load_dotenv

load_dotenv()
RESEND_API_KEY = os.getenv("RESEND_API_KEY", "").strip()
RESEND_FROM_EMAIL = os.getenv(
    "RESEND_FROM_EMAIL",
    "onboarding@resend.dev"
).strip()

APP_NAME = os.getenv("APP_NAME", "CREOVANTA API")
JWT_SECRET = os.getenv("JWT_SECRET", "dev-only-change-this-secret-before-deployment")
ALGORITHM = "HS256"
TOKEN_MINUTES = int(os.getenv("ACCESS_TOKEN_MINUTES", "720"))
app = FastAPI(title=APP_NAME, version="1.0.0", description="CREOVANTA AI Content Creator Marketplace API")
# Local development origins plus any frontend origins configured in the environment.
# In Render, set FRONTEND_ORIGIN to your deployed frontend URL.
default_origins = [
    "http://localhost:5173",
    "http://127.0.0.1:5173",
    "http://localhost:5174",
    "http://127.0.0.1:5174",
    "http://10.55.184.34:5173",
    "http://10.55.184.34:5174",
    "http://10.119.188.34:5173",
    "http://10.119.188.34:5174",
    "https://creovanta-corrected-project-zip-2-if8e.onrender.com",
    "https://collabcrew.onrender.com",
]
configured_origins = [
    origin.strip().rstrip("/")
    for origin in os.getenv("FRONTEND_ORIGIN", "").split(",")
    if origin.strip()
]
origins = list(dict.fromkeys(default_origins + configured_origins))

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
security = HTTPBearer(auto_error=False)

# Demo-only passwords are hashed with PBKDF2; never store plaintext passwords.
def hash_password(password: str) -> str:
    salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 180000).hex()
    return f"pbkdf2_sha256${salt}${digest}"

def verify_password(password: str, encoded: str) -> bool:
    try:
        _, salt, digest = encoded.split("$", 2)
        candidate = hashlib.pbkdf2_hmac("sha256", password.encode(), salt.encode(), 180000).hex()
        return secrets.compare_digest(candidate, digest)
    except Exception:
        return False

def send_verification_email(to_email: str, otp: str):
    if not RESEND_API_KEY:
        raise HTTPException(
            status_code=503,
            detail="Resend is not configured. Check the backend .env file."
        )

    try:
        response = requests.post(
            "https://api.resend.com/emails",
            headers={
                "Authorization": f"Bearer {RESEND_API_KEY}",
                "Content-Type": "application/json",
            },
            json={
                "from": RESEND_FROM_EMAIL,
                "to": [to_email],
                "subject": "Verify your CREOVANTA email",
                "html": f"""
                    <div style="font-family:Arial,sans-serif">
                      <h2>Welcome to CREOVANTA!</h2>
                      <p>Your email verification code is:</p>
                      <h1 style="letter-spacing:6px">{otp}</h1>
                      <p>This code expires in 10 minutes.</p>
                      <p>If you didn't request this code, ignore this email.</p>
                    </div>
                """,
            },
            timeout=15,
        )
    except requests.RequestException:
        raise HTTPException(
            status_code=502,
            detail="Could not contact the email service. Please try again."
        )

    if not response.ok:
        print("Resend error:", response.status_code, response.text)
        raise HTTPException(
            status_code=502,
            detail="Email could not be sent. Check the Resend sender and domain settings."
        )

def _b64(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).rstrip(b"=").decode()

def create_token(user: User) -> str:
    now = datetime.utcnow()
    header = _b64(json.dumps({"alg": "HS256", "typ": "JWT"}, separators=(",", ":")).encode())
    payload = _b64(json.dumps({"sub": str(user.id), "role": user.role, "iat": int(now.timestamp()), "exp": int((now + timedelta(minutes=TOKEN_MINUTES)).timestamp())}, separators=(",", ":")).encode())
    signing_input = f"{header}.{payload}"
    signature = _b64(hmac.new(JWT_SECRET.encode(), signing_input.encode(), hashlib.sha256).digest())
    return f"{signing_input}.{signature}"

def decode_token(token: str) -> dict:
    try:
        header, payload, signature = token.split(".")
        signing_input = f"{header}.{payload}"
        expected = _b64(hmac.new(JWT_SECRET.encode(), signing_input.encode(), hashlib.sha256).digest())
        if not hmac.compare_digest(signature, expected):
            raise ValueError("Invalid signature")
        decoded = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
        if decoded.get("exp", 0) <= int(datetime.utcnow().timestamp()):
            raise ValueError("Expired token")
        return decoded
    except Exception as exc:
        raise ValueError("Invalid or expired token") from exc

def user_out(user: User):
    return {
        "id": user.id,
        "name": user.name,
        "email": user.email,
        "role": user.role,
        "avatar_url": user.avatar_url,
        "email_verified": bool(getattr(user, "email_verified", False)),
    }

def profile_out(p: CreatorProfile, include_private=False):
    user = p.user
    projects = [project_out(x, include_private=include_private) for x in p.projects if include_private or x.visibility == "Public"]
    return {"id": p.id, "user_id": p.user_id, "name": user.name, "email": user.email if include_private else None,
        "username": p.username, "biography": p.biography, "specialization": p.specialization,
        "experience_level": p.experience_level, "availability": p.availability, "rate_min": p.rate_min, "rate_max": p.rate_max,
        "skills": p.skills or [], "tools": p.tools or [], "content_types": p.content_types or [], "formats": p.formats or [],
        "aspect_ratios": p.aspect_ratios or [], "languages": p.languages or [], "workflow": p.workflow,
        "verification_status": p.verification_status, "portfolio_count": len(projects), "projects": projects,
        "avatar_url": user.avatar_url, "is_demo": user.email.endswith("@demo.creovanta.app")}

def project_out(p: PortfolioProject, include_private=False):
    return {"id": p.id, "creator_id": p.creator_id, "title": p.title, "description": p.description,
        "content_type": p.content_type, "visual_style": p.visual_style, "preview_url": p.preview_url,
        "tools": p.tools or [], "skills": p.skills or [], "workflow": p.workflow,
        "prompt_notes": p.prompt_notes if (include_private or not p.prompt_private) else "",
        "prompt_private": p.prompt_private, "aspect_ratio": p.aspect_ratio, "output_format": p.output_format,
        "commercial_use": p.commercial_use, "license_type": p.license_type, "visibility": p.visibility,
        "is_demo": p.is_demo, "created_at": p.created_at.isoformat() if p.created_at else None}

def brief_out(b: CreativeBrief, include_owner=False):
    return {"id": b.id, "brand_user_id": b.brand_user_id, "brand_name": b.brand_name, "title": b.title,
        "description": b.description, "industry": b.industry, "objective": b.objective, "target_audience": b.target_audience,
        "content_type": b.content_type, "visual_style": b.visual_style, "platform": b.platform, "aspect_ratio": b.aspect_ratio,
        "output_format": b.output_format, "budget_min": b.budget_min, "budget_max": b.budget_max,
        "deadline": b.deadline.isoformat() if b.deadline else None, "deliverables": b.deliverables,
        "required_skills": b.required_skills or [], "required_tools": b.required_tools or [],
        "commercial_use": b.commercial_use, "license_type": b.license_type, "usage_rights": b.usage_rights,
        "ai_disclosure": b.ai_disclosure, "revision_rounds": b.revision_rounds, "status": b.status,
        "is_demo": b.is_demo, "created_at": b.created_at.isoformat() if b.created_at else None}

def current_user(credentials: HTTPAuthorizationCredentials = Depends(security), db: Session = Depends(get_db)) -> User:
    if not credentials:
        raise HTTPException(401, "Sign in to continue")
    try:
        payload = decode_token(credentials.credentials)
        user_id = int(payload.get("sub", "0"))
    except (ValueError, TypeError):
        raise HTTPException(401, "Invalid or expired session")
    user = db.get(User, user_id)
    if not user:
        raise HTTPException(401, "Account not found")
    return user

def require_role(user: User, *roles):
    if user.role not in roles:
        raise HTTPException(403, f"This action requires one of these roles: {', '.join(roles)}")

def get_profile_for_user(db: Session, user_id: int):
    return db.query(CreatorProfile).options(joinedload(CreatorProfile.user), joinedload(CreatorProfile.projects)).filter_by(user_id=user_id).first()

def seed_data(db: Session):
    if db.query(User).count() > 0:
        return
    brand = User(name="Maya Chen", email="brand@demo.creovanta.app", password_hash=hash_password("DemoPass123!"), role="brand", avatar_url="MC")
    admin = User(name="Alex Morgan", email="admin@demo.creovanta.app", password_hash=hash_password("DemoPass123!"), role="admin", avatar_url="AM")
    db.add_all([brand, admin]); db.flush()
    names = ["Aarav Mehta", "Zara Kapoor", "Ishaan Rao", "Mira Shah", "Dev Patel", "Anika Bose", "Kabir Nair", "Rhea Iyer", "Neil Das", "Tara Menon", "Arjun Sen", "Kiara Reddy"]
    specializations = ["AI Filmmaker", "Generative Artist", "AI Animator", "Product Ads", "Social Video", "3D Artist"]
    tools_list = [["Runway", "Midjourney"], ["Midjourney", "Photoshop AI"], ["Pika", "Blender"], ["Stable Diffusion", "ComfyUI"], ["Sora", "CapCut"], ["DALL·E", "After Effects"]]
    skills_list = [["Prompt Design", "Cinematic Lighting", "Storyboarding"], ["Generative Art", "Color Grading"], ["Motion Design", "3D Animation"], ["ComfyUI", "Image Editing"], ["Ad Creative", "Video Editing"], ["Concept Art", "Visual Storytelling"]]
    for i, name in enumerate(names):
        email = f"creator{i+1}@demo.creovanta.app"
        user = User(name=name, email=email, password_hash=hash_password("DemoPass123!"), role="creator", avatar_url="".join(x[0] for x in name.split()))
        p = CreatorProfile(username=name.lower().replace(" ", ""), biography=f"I create distinctive AI-native visuals for brands, storytellers, and digital-first campaigns. Focused on craft, transparent workflows, and reliable delivery.", specialization=specializations[i % len(specializations)], experience_level=["Intermediate", "Advanced", "Expert"][i % 3], availability="Available" if i % 4 else "Booked until next month", rate_min=5000 + i*500, rate_max=12000 + i*700, skills=skills_list[i % len(skills_list)], tools=tools_list[i % len(tools_list)], content_types=[["AI Video", "Advertisement"], ["AI Image", "Experimental Art"], ["Animation", "AI Video"]][i % 3], formats=["MP4", "PNG", "WEBP"], aspect_ratios=["9:16", "16:9", "1:1"], languages=["English", "Hindi"] if i % 2 else ["English", "Telugu"], workflow="Brief discovery → visual references → prompt iterations → generation → edit and grade → QA and delivery", verification_status=["Self-reported", "Pending review", "Workflow reviewed"][i % 3])
        user.creator_profile = p
        db.add(user); db.flush()
        for j in range(2):
            p.projects.append(PortfolioProject(title=["Neon Future Product Film", "Dreamscape Campaign", "Orbital Brand Identity", "Liquid Chrome Study"][((i+j) % 4)], description="A demo portfolio concept exploring a bold visual direction for a modern digital audience.", content_type=["AI Video", "AI Image", "Animation", "Advertisement"][(i+j) % 4], visual_style=["Cinematic", "Futuristic", "3D", "Editorial"][(i+j) % 4], preview_url="", tools=tools_list[(i+j) % len(tools_list)], skills=skills_list[(i+j) % len(skills_list)], workflow="Concept → reference board → prompt iteration → generation → post-production → export", prompt_notes="Demo prompt notes: cinematic lighting, deliberate composition, premium color grade.", prompt_private=(j == 0), aspect_ratio=["9:16", "16:9", "1:1"][j % 3], output_format="MP4" if j == 0 else "PNG", commercial_use=(j == 1), license_type="Demo only — confirm commercial terms", visibility="Public", is_demo=True))
    brief_data = [
        ("Futuristic skincare launch", "Create a premium 20-second vertical product film with clean highlights and sculptural lighting.", "Beauty", "AI Video", "Cinematic", 18000, 30000, ["Product Lighting", "Video Editing"]),
        ("Next-gen sneaker campaign", "Develop a set of energetic AI visuals and one motion teaser for a new sneaker drop.", "Fashion", "Advertisement", "3D", 22000, 40000, ["3D Animation", "Generative Art"]),
        ("Sustainable living explainer", "Produce a friendly animated explainer for a sustainability-focused consumer brand.", "Environment", "Animation", "Minimal", 12000, 25000, ["Motion Design", "Storyboarding"]),
        ("Music festival visual loop", "Create a seamless looping visual for stage screens and social promotion.", "Entertainment", "Animation", "Experimental", 15000, 28000, ["Motion Design", "Visual Storytelling"]),
        ("Fintech app launch reel", "Create a sharp 15-second social reel introducing a new personal finance app.", "Finance", "Social Video", "Futuristic", 14000, 26000, ["Ad Creative", "Video Editing"]),
        ("Travel destination mood film", "Craft an atmospheric short film that sells the feeling of an undiscovered destination.", "Travel", "AI Video", "Editorial", 20000, 35000, ["Cinematic Lighting", "Color Grading"]),
        ("Gaming character reveal", "Produce a stylized character reveal with a strong silhouette and dramatic motion.", "Gaming", "AI Video", "Anime", 16000, 32000, ["Concept Art", "Prompt Design"]),
        ("Smart home product visuals", "Create a coherent series of photoreal product visuals for a smart-home launch.", "Technology", "AI Image", "Realistic", 10000, 22000, ["Product Lighting", "Generative Art"])]
    for i, (title, desc, industry, ctype, style, bmin, bmax, skills) in enumerate(brief_data):
        db.add(CreativeBrief(brand_user_id=brand.id, brand_name=["Nova Beauty", "Stride Labs", "Greenloop", "Pulse Festival"][i%4], title=title, description=desc, industry=industry, objective="Launch campaign", target_audience="Digital-first audience, ages 18–35", content_type=ctype, visual_style=style, platform="Instagram" if i % 2 == 0 else "YouTube", aspect_ratio="9:16" if i % 2 == 0 else "16:9", output_format="MP4" if "Video" in ctype or ctype == "Animation" else "PNG", budget_min=bmin, budget_max=bmax, deadline=date.today()+timedelta(days=10+i*2), deliverables=2 if i % 3 else 3, required_skills=skills, required_tools=["Runway", "Midjourney"], commercial_use=True, license_type="Non-exclusive campaign license", usage_rights="Digital campaign use; final scope to be agreed in writing", ai_disclosure="Disclose AI-generated content", revision_rounds=2, status="Published" if i < 6 else "Draft", is_demo=True))
    db.commit()

@app.on_event("startup")
def startup():
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        seed_data(db)
        # Demo accounts are pre-verified; this also updates demo accounts in existing databases.
        db.query(User).filter(User.email.ilike("%@demo.creovanta.app")).update(
            {User.email_verified: True, User.otp_hash: None, User.otp_expires_at: None},
            synchronize_session=False,
        )
        db.commit()

@app.get("/")
def root():
    # Never expose demo passwords or account credentials in a public API response.
    return {"name": "CREOVANTA API", "status": "online", "docs": "/docs"}

@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    return {"status": "healthy", "database": "connected", "creators": db.query(CreatorProfile).count(), "briefs": db.query(CreativeBrief).count()}

@app.post("/api/auth/register")
def register(payload: RegisterIn, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    name = payload.name.strip()

    if "@" not in email or "." not in email.split("@")[-1]:
        raise HTTPException(422, "Enter a valid email address")

    if not name:
        raise HTTPException(422, "Enter your full name")

    if payload.role not in ["creator", "brand"]:
        raise HTTPException(422, "Role must be creator or brand")

    if db.query(User).filter_by(email=email).first():
        raise HTTPException(
            409, "An account with this email already exists"
        )

    # Generate a six-digit OTP and store only its hash.
    otp = f"{secrets.randbelow(1_000_000):06d}"
    otp_hash = hashlib.sha256(
        f"{email}:{otp}".encode()
    ).hexdigest()

    # Send the email before creating the account.
    send_verification_email(email, otp)

    user = User(
        name=name,
        email=email,
        password_hash=hash_password(payload.password),
        role=payload.role,
        avatar_url="".join(
            part[0] for part in name.split()[:2]
        ).upper(),
        email_verified=False,
        otp_hash=otp_hash,
        otp_expires_at=datetime.utcnow() + timedelta(minutes=10),
    )

    try:
        db.add(user)
        db.flush()

        if user.role == "creator":
            username = "".join(
                ch for ch in name.lower().replace(" ", "")
                if ch.isalnum()
            ) + str(user.id)

            user.creator_profile = CreatorProfile(
                username=username,
                biography="Add a short introduction to your creative practice.",
                specialization="AI Video",
                skills=[],
                tools=[],
                content_types=["AI Video"],
                formats=["MP4", "PNG"],
                aspect_ratios=["9:16", "16:9"],
                languages=["English"],
                workflow="Describe your workflow to help brands understand your process.",
            )

        db.commit()
        db.refresh(user)

    except Exception:
        db.rollback()
        raise

    return {
        "message": "Verification code sent to your email.",
        "email": email,
        "requires_verification": True,
    }
@app.post("/api/auth/resend-otp")
def resend_otp(payload: dict, db: Session = Depends(get_db)):
    email = str(payload.get("email", "")).strip().lower()

    if not email or "@" not in email:
        raise HTTPException(422, "Enter a valid email address")

    user = db.query(User).filter_by(email=email).first()

    if not user:
        raise HTTPException(404, "Account not found. Please register first.")

    if user.email_verified:
        raise HTTPException(400, "Email is already verified. Please sign in.")

    otp = f"{secrets.randbelow(1_000_000):06d}"
    otp_hash = hashlib.sha256(f"{email}:{otp}".encode()).hexdigest()

    # Send the email before updating the stored OTP.
    send_verification_email(email, otp)

    user.otp_hash = otp_hash
    user.otp_expires_at = datetime.utcnow() + timedelta(minutes=10)
    db.commit()

    return {
        "message": "Verification code sent to your email.",
        "email": email,
        "requires_verification": True,
    }

@app.post("/api/auth/verify-otp")
def verify_otp(payload: dict, db: Session = Depends(get_db)):
    email = str(payload.get("email", "")).strip().lower()
    otp = str(payload.get("otp", "")).strip()

    if len(otp) != 6 or not otp.isdigit():
        raise HTTPException(422, "Enter the six-digit verification code")

    user = db.query(User).filter_by(email=email).first()

    if not user:
        raise HTTPException(404, "Account not found. Please register again.")

    if user.email_verified:
        raise HTTPException(400, "Email is already verified. Please sign in.")

    if not user.otp_hash or not user.otp_expires_at:
        raise HTTPException(400, "No active verification code.")

    if datetime.utcnow() > user.otp_expires_at:
        user.otp_hash = None
        user.otp_expires_at = None
        db.commit()
        raise HTTPException(400, "Code expired. Please register again.")

    candidate_hash = hashlib.sha256(
        f"{email}:{otp}".encode()
    ).hexdigest()

    if not hmac.compare_digest(candidate_hash, user.otp_hash):
        raise HTTPException(400, "Incorrect verification code.")

    user.email_verified = True
    user.otp_hash = None
    user.otp_expires_at = None

    db.commit()
    db.refresh(user)

    return {
        "access_token": create_token(user),
        "token_type": "bearer",
        "user": user_out(user),
    }

@app.post("/api/auth/login")
def login(payload: LoginIn, db: Session = Depends(get_db)):
    email = payload.email.strip().lower()
    user = db.query(User).filter_by(email=email).first()

    if not user or not verify_password(
        payload.password, user.password_hash
    ):
        raise HTTPException(401, "Email or password is incorrect")

    is_demo = email.endswith("@demo.creovanta.app")

    if not is_demo and not getattr(user, "email_verified", False):
        raise HTTPException(
            403,
            "Please verify your email before signing in."
        )

    return {
        "access_token": create_token(user),
        "token_type": "bearer",
        "user": user_out(user),
    }


@app.get("/api/auth/me")
def me(user: User = Depends(current_user)):
    return user_out(user)


@app.get("/api/admin/users.xlsx")
def export_users_excel(user: User = Depends(current_user), db: Session = Depends(get_db)):
    """Download a fresh Excel snapshot of all users. Admin-only; never exports passwords or hashes."""
    require_role(user, "admin")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter

    users = db.query(User).order_by(User.id.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "All Accounts"
    ws.append([
        "User ID",
        "Name",
        "Email",
        "Role",
        "Account Type",
        "Email Verified",
        "Registered Date",
    ])

    for account in users:
        is_demo = account.email.lower().endswith("@demo.creovanta.app")
        verified = bool(getattr(account, "email_verified", False)) or is_demo
        ws.append([
            account.id,
            account.name,
            account.email,
            account.role.title(),
            "Demo" if is_demo else "Real",
            "Verified" if verified else "Not verified",
            account.created_at.strftime("%Y-%m-%d %H:%M:%S")
            if account.created_at else "",
        ])

    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, width in enumerate([12, 28, 40, 14, 16, 18, 23], start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    if users:
        table = Table(displayName="CreovantaUsers", ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"creovanta_accounts_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/api/creators")
def list_creators(search: str = "", skills: Optional[str] = None, tools: Optional[str] = None, specialization: Optional[str] = None, content_type: Optional[str] = None, aspect_ratio: Optional[str] = None, availability: Optional[str] = None, experience_level: Optional[str] = None, verified: Optional[str] = None, sort: str = "relevant", page: int = Query(1, ge=1), page_size: int = Query(12, ge=1, le=50), db: Session = Depends(get_db)):
    q = db.query(CreatorProfile).options(joinedload(CreatorProfile.user), joinedload(CreatorProfile.projects))
    q = q.join(User)
    items = q.all()
    def matches(p):
        def has_all(field, val): return not val or all(any(v.lower() in str(x).lower() for x in (getattr(p, field) or [])) for v in val.split(",") if v)
        if not has_all("skills", skills) or not has_all("tools", tools) or not has_all("content_types", content_type) or not has_all("aspect_ratios", aspect_ratio): return False
        if specialization and specialization.lower() not in p.specialization.lower(): return False
        if availability and availability.lower() not in p.availability.lower(): return False
        if experience_level and experience_level.lower() not in p.experience_level.lower(): return False
        if verified and verified != "Any" and p.verification_status.lower() != verified.lower(): return False
        if search:
            text = " ".join([p.user.name, p.biography, p.specialization, *(p.skills or []), *(p.tools or []), *(p.content_types or []), *(x.title for x in p.projects if x.visibility == "Public")]).lower()
            if search.lower() not in text: return False
        return True
    items = [p for p in items if matches(p)]
    if sort == "projects": items.sort(key=lambda p: len([x for x in p.projects if x.visibility == "Public"]), reverse=True)
    elif sort == "rate_low": items.sort(key=lambda p: p.rate_min if p.rate_min is not None else 10**9)
    elif sort == "rate_high": items.sort(key=lambda p: p.rate_max if p.rate_max is not None else -1, reverse=True)
    elif sort == "verified": items.sort(key=lambda p: (p.verification_status == "Workflow reviewed", p.verification_status == "Past work reviewed"), reverse=True)
    elif sort == "recent": items.sort(key=lambda p: p.user.created_at, reverse=True)
    total = len(items); start = (page-1)*page_size
    return {"items": [profile_out(p) for p in items[start:start+page_size]], "total": total, "page": page, "page_size": page_size, "pages": (total+page_size-1)//page_size}

@app.get("/api/creators/{creator_id}")
def get_creator(creator_id: int, db: Session = Depends(get_db)):
    p = db.query(CreatorProfile).options(joinedload(CreatorProfile.user), joinedload(CreatorProfile.projects)).filter_by(id=creator_id).first()
    if not p: raise HTTPException(404, "Creator not found")
    return profile_out(p)

@app.get("/api/creators/me/profile")
def my_profile(user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "creator", "admin")
    p = get_profile_for_user(db, user.id)
    if not p: raise HTTPException(404, "Creator profile not found")
    return profile_out(p, include_private=True)

@app.put("/api/creators/me/profile")
def update_profile(payload: ProfileIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "creator", "admin")
    p = get_profile_for_user(db, user.id)
    if not p: raise HTTPException(404, "Creator profile not found")
    data = payload.model_dump(exclude_unset=True)
    if "name" in data: user.name = data.pop("name")
    for k, v in data.items():
        if hasattr(p, k): setattr(p, k, v)
    db.commit(); db.refresh(p)
    return profile_out(p, include_private=True)

@app.post("/api/portfolio")
def create_portfolio(payload: PortfolioIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "creator", "admin")
    p = get_profile_for_user(db, user.id)
    if not p: raise HTTPException(404, "Creator profile not found")
    project = PortfolioProject(creator_id=p.id, **payload.model_dump())
    db.add(project); db.commit(); db.refresh(project)
    return project_out(project, include_private=True)

@app.get("/api/portfolio/{project_id}")
def get_portfolio(project_id: int, db: Session = Depends(get_db)):
    project = db.get(PortfolioProject, project_id)
    if not project or project.visibility != "Public": raise HTTPException(404, "Portfolio project not found")
    return project_out(project)

@app.put("/api/portfolio/{project_id}")
def update_portfolio(project_id: int, payload: PortfolioIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = db.get(PortfolioProject, project_id)
    if not project: raise HTTPException(404, "Portfolio project not found")
    p = db.get(CreatorProfile, project.creator_id)
    if user.role != "admin" and (not p or p.user_id != user.id): raise HTTPException(403, "You can only edit your own portfolio projects")
    for k, v in payload.model_dump().items(): setattr(project, k, v)
    db.commit(); db.refresh(project)
    return project_out(project, include_private=True)

@app.delete("/api/portfolio/{project_id}")
def delete_portfolio(project_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    project = db.get(PortfolioProject, project_id)
    if not project: raise HTTPException(404, "Portfolio project not found")
    p = db.get(CreatorProfile, project.creator_id)
    if user.role != "admin" and (not p or p.user_id != user.id): raise HTTPException(403, "You can only delete your own portfolio projects")
    db.delete(project); db.commit(); return {"message": "Portfolio project deleted"}

@app.get("/api/briefs")
def list_briefs(status: Optional[str] = None, search: str = "", mine: bool = False, user: Optional[User] = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(CreativeBrief)
    if status: q = q.filter(CreativeBrief.status == status)
    if search: q = q.filter(or_(CreativeBrief.title.ilike(f"%{search}%"), CreativeBrief.description.ilike(f"%{search}%"), CreativeBrief.industry.ilike(f"%{search}%")))
    if mine and user: q = q.filter(CreativeBrief.brand_user_id == user.id)
    return [brief_out(b) for b in q.order_by(desc(CreativeBrief.created_at)).all()]

@app.post("/api/briefs")
def create_brief(payload: BriefIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "brand", "admin")
    data = payload.model_dump()
    if data["budget_min"] is not None and data["budget_max"] is not None and data["budget_min"] > data["budget_max"]: raise HTTPException(422, "Minimum budget cannot exceed maximum budget")
    if data["status"] not in ["Draft", "Published"]: data["status"] = "Draft"
    b = CreativeBrief(brand_user_id=user.id, **data)
    db.add(b); db.commit(); db.refresh(b); return brief_out(b)

@app.get("/api/briefs/{brief_id}")
def get_brief(brief_id: int, db: Session = Depends(get_db)):
    b = db.get(CreativeBrief, brief_id)
    if not b: raise HTTPException(404, "Creative brief not found")
    return brief_out(b)

@app.put("/api/briefs/{brief_id}")
def update_brief(brief_id: int, payload: BriefIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    b = db.get(CreativeBrief, brief_id)
    if not b: raise HTTPException(404, "Creative brief not found")
    if user.role != "admin" and b.brand_user_id != user.id: raise HTTPException(403, "You can only edit your own briefs")
    data = payload.model_dump()
    if data["budget_min"] is not None and data["budget_max"] is not None and data["budget_min"] > data["budget_max"]: raise HTTPException(422, "Minimum budget cannot exceed maximum budget")
    for k, v in data.items(): setattr(b, k, v)
    db.commit(); db.refresh(b); return brief_out(b)

@app.post("/api/briefs/{brief_id}/publish")
def publish_brief(brief_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    b = db.get(CreativeBrief, brief_id)
    if not b: raise HTTPException(404, "Creative brief not found")
    if user.role != "admin" and b.brand_user_id != user.id: raise HTTPException(403, "You can only publish your own briefs")
    b.status = "Published"; db.commit(); return brief_out(b)

@app.delete("/api/briefs/{brief_id}")
def delete_brief(brief_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    b = db.get(CreativeBrief, brief_id)
    if not b: raise HTTPException(404, "Creative brief not found")
    if user.role != "admin" and b.brand_user_id != user.id: raise HTTPException(403, "You can only delete your own briefs")
    db.delete(b); db.commit(); return {"message": "Brief deleted"}

@app.post("/api/briefs/{brief_id}/applications")
def apply_to_brief(brief_id: int, payload: ApplicationIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "creator", "admin")
    b = db.get(CreativeBrief, brief_id)
    if not b or b.status != "Published": raise HTTPException(404, "Published brief not found")
    p = get_profile_for_user(db, user.id)
    if not p: raise HTTPException(400, "Complete your creator profile first")
    existing = db.query(Application).filter_by(brief_id=brief_id, creator_id=p.id).first()
    if existing: raise HTTPException(409, "You have already applied to this brief")
    a = Application(brief_id=brief_id, creator_id=p.id, **payload.model_dump())
    db.add(a); db.commit(); db.refresh(a)
    return {"id": a.id, "brief_id": a.brief_id, "creator_id": a.creator_id, "proposal": a.proposal, "proposed_budget": a.proposed_budget, "status": a.status}

@app.get("/api/briefs/{brief_id}/applications")
def list_applications(brief_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)):
    b = db.get(CreativeBrief, brief_id)
    if not b: raise HTTPException(404, "Brief not found")
    if user.role != "admin" and b.brand_user_id != user.id: raise HTTPException(403, "Only the brief owner can review applications")
    rows = db.query(Application).filter_by(brief_id=brief_id).all()
    result = []
    for a in rows:
        p = db.query(CreatorProfile).options(joinedload(CreatorProfile.user)).filter_by(id=a.creator_id).first()
        result.append({"id": a.id, "brief_id": a.brief_id, "creator_id": a.creator_id, "creator_name": p.user.name if p else "Creator", "username": p.username if p else "", "proposal": a.proposal, "proposed_budget": a.proposed_budget, "status": a.status, "created_at": a.created_at.isoformat()})
    return result

@app.patch("/api/applications/{application_id}")
def update_application(application_id: int, payload: StatusIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    a = db.get(Application, application_id)
    if not a: raise HTTPException(404, "Application not found")
    b = db.get(CreativeBrief, a.brief_id)
    if user.role != "admin" and (not b or b.brand_user_id != user.id): raise HTTPException(403, "Only the brief owner can update applications")
    if payload.status not in ["Shortlisted", "Accepted", "Rejected", "In Discussion"]: raise HTTPException(422, "Unsupported application status")
    a.status = payload.status
    if payload.status == "Accepted":
        exists = db.query(Collaboration).filter_by(brief_id=a.brief_id, creator_id=a.creator_id).first()
        if not exists: db.add(Collaboration(brief_id=a.brief_id, brand_user_id=b.brand_user_id, creator_id=a.creator_id, status="Accepted", activity_log=[{"at": datetime.utcnow().isoformat(), "event": "Application accepted"}]))
    db.commit(); return {"id": a.id, "status": a.status}

@app.get("/api/collaborations")
def list_collaborations(user: User = Depends(current_user), db: Session = Depends(get_db)):
    q = db.query(Collaboration)
    if user.role == "creator":
        p = get_profile_for_user(db, user.id); q = q.filter(Collaboration.creator_id == (p.id if p else -1))
    elif user.role != "admin": q = q.filter(Collaboration.brand_user_id == user.id)
    result = []
    for c in q.order_by(desc(Collaboration.created_at)).all():
        b = db.get(CreativeBrief, c.brief_id); p = db.query(CreatorProfile).options(joinedload(CreatorProfile.user)).filter_by(id=c.creator_id).first()
        result.append({"id": c.id, "brief_id": c.brief_id, "brief_title": b.title if b else "Deleted brief", "creator_id": c.creator_id, "creator_name": p.user.name if p else "Creator", "status": c.status, "activity_log": c.activity_log or [], "created_at": c.created_at.isoformat()})
    return result

@app.patch("/api/collaborations/{collaboration_id}")
def update_collaboration(collaboration_id: int, payload: StatusIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    c = db.get(Collaboration, collaboration_id)
    if not c: raise HTTPException(404, "Collaboration not found")
    p = db.get(CreatorProfile, c.creator_id)
    if user.role != "admin" and user.id not in [c.brand_user_id, p.user_id if p else -1]: raise HTTPException(403, "You are not part of this collaboration")
    allowed = ["Accepted", "In Discussion", "In Progress", "Submitted", "Revision Requested", "Approved", "Completed", "Cancelled"]
    if payload.status not in allowed: raise HTTPException(422, "Unsupported collaboration status")
    c.status = payload.status; log = c.activity_log or []; log.append({"at": datetime.utcnow().isoformat(), "event": f"Status changed to {payload.status}", "by": user.name}); c.activity_log = log
    db.commit(); return {"id": c.id, "status": c.status, "activity_log": c.activity_log}

@app.post("/api/ai/brief-builder")
def ai_brief_builder(payload: dict, user: User = Depends(current_user)):
    idea = str(payload.get("idea", "")).strip()
    if len(idea) < 10: raise HTTPException(422, "Describe your campaign idea in at least 10 characters")
    text = idea.lower()
    ctype = "Animation" if any(x in text for x in ["animate", "animation", "animated"]) else "AI Image" if any(x in text for x in ["image", "poster", "still", "visuals"]) else "AI Video"
    style = "Futuristic" if any(x in text for x in ["future", "futuristic", "tech", "sci-fi"]) else "Realistic" if any(x in text for x in ["realistic", "photoreal", "product"] ) else "Cinematic"
    platform = "Instagram Reels" if any(x in text for x in ["instagram", "reel", "vertical", "tiktok"]) else "YouTube" if "youtube" in text else "Digital campaign"
    ratio = "9:16" if platform == "Instagram Reels" else "16:9" if platform == "YouTube" else "1:1"
    title = " ".join(idea.split()[:6]).strip(".,!? ").title()
    if len(title) < 3: title = "Creative Campaign Concept"
    return {"engine": "CREOVANTA fallback rules engine", "is_fallback": True, "notice": "Suggestions only. Confirm budget, deadline, commercial rights, and final requirements before publishing.", "brief": {"title": title, "description": idea, "industry": "Technology" if style == "Futuristic" else "Lifestyle", "objective": "Product awareness", "target_audience": "To be confirmed", "content_type": ctype, "visual_style": style, "platform": platform, "aspect_ratio": ratio, "output_format": "MP4" if ctype != "AI Image" else "PNG", "deliverables": 1, "required_skills": ["Prompt Design", "Visual Storytelling", "Post-production"], "required_tools": ["Runway", "Midjourney"], "budget_min": None, "budget_max": None, "deadline": None, "commercial_use": False, "license_type": "Confirm with creator", "usage_rights": "Requires confirmation", "ai_disclosure": "Confirm disclosure requirements", "revision_rounds": 2, "status": "Draft"}, "missing_details": ["Budget", "Deadline", "Commercial usage rights", "Territory and duration of license"]}

@app.post("/api/verification/evidence")
def submit_evidence(payload: EvidenceIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "creator", "admin")
    p = get_profile_for_user(db, user.id)
    if not p: raise HTTPException(404, "Creator profile not found")
    ev = VerificationEvidence(creator_id=p.id, evidence_type=payload.evidence_type, evidence_reference=payload.evidence_reference, review_status="Pending")
    db.add(ev); db.commit(); db.refresh(ev)
    return {"id": ev.id, "review_status": ev.review_status, "message": "Evidence submitted for admin review"}

@app.get("/api/verification/pending")
def pending_evidence(user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "admin")
    rows = db.query(VerificationEvidence).filter_by(review_status="Pending").all()
    return [{"id": e.id, "creator_id": e.creator_id, "creator_name": (db.query(CreatorProfile).options(joinedload(CreatorProfile.user)).filter_by(id=e.creator_id).first().user.name), "evidence_type": e.evidence_type, "evidence_reference": e.evidence_reference, "review_status": e.review_status} for e in rows]

@app.patch("/api/verification/{evidence_id}")
def review_evidence(evidence_id: int, payload: StatusIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    require_role(user, "admin")
    ev = db.get(VerificationEvidence, evidence_id)
    if not ev: raise HTTPException(404, "Evidence not found")
    if payload.status not in ["Verified", "Rejected"]: raise HTTPException(422, "Status must be Verified or Rejected")
    ev.review_status = payload.status; ev.reviewed_by = user.id
    if payload.status == "Verified":
        p = db.get(CreatorProfile, ev.creator_id); p.verification_status = "Workflow reviewed" if "workflow" in ev.evidence_type.lower() else "Past work reviewed"
    db.commit(); return {"id": ev.id, "review_status": ev.review_status}

@app.get("/api/analytics/dashboard")
def analytics(db: Session = Depends(get_db)):
    return {"demo_data": False, "total_creators": db.query(CreatorProfile).count(), "portfolio_projects": db.query(PortfolioProject).count(), "active_briefs": db.query(CreativeBrief).filter(CreativeBrief.status == "Published").count(), "active_collaborations": db.query(Collaboration).filter(Collaboration.status.in_(["Accepted", "In Progress", "In Discussion"])).count(), "applications": db.query(Application).count(), "monthly_activity": [], "popular_formats": []}
